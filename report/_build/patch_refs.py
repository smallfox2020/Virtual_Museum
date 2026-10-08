# -*- coding: utf-8 -*-
"""在用户自己改过的 开题报告.docx / 任务书.docx 上，只替换参考文献与文献综述里的引用标注。
   读 report/*.docx，写到 _build/out/*.docx（磁盘上的原件被 Word 锁住时也能跑）。"""
import os, sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # report/
OUT = os.environ.get('OUT_DIR') or os.path.join(HERE, '_build', 'out')
os.makedirs(OUT, exist_ok=True)

REFS = [
    '[1]孙丹妮.欧美虚拟博物馆理论与馆建述略[J].美术,2020,(06):20-26.',
    '[2]向辉,孟祥旭,杨承磊.山东大学考古数字博物馆设计与实现[J].系统仿真学报,2003,15(03):319-321.',
    '[3]曾定浩,卢威,贝佳,等.一个基于XML的虚拟博物馆场景描述语言[J].系统仿真学报,2006,18(09):2492-2496.',
    '[4]曹彤.虚拟博物馆的三维场景构造及交互漫游实现[J].计算机工程与设计,2007,28(24):6006-6007,6011.',
    '[5]李赟,刘一松,陈继明.基于WebGL的三维虚拟博物馆碰撞检测[J].软件导刊,2017,16(05):124-128.',
    '[6]任国栋,陈林华,陶学锋,等.基于Unity3D的虚拟博物馆信息可视化系统[J].计算机系统应用,2013,22(09):86-90,59.',
    '[7]马自萍,郭贝贝,李海东,等.Pano2VR的宁夏虚拟博物馆全景漫游实现[J].现代电子技术,2021,44(08):149-153.'
    'DOI:10.16652/j.issn.1004-373x.2021.08.033.',
    '[8]裴卉宁,温志强,黄雪芹,等.融合视觉认知特征的虚拟博物馆界面布局美度评价方法[J].图学学报,2023,44(02):389-398.'
    'DOI:10.11996/JG.j.2095-302X.2023020389.',
    '[9]王红,刘素仁.沉浸与叙事:新媒体影像技术下的博物馆文化沉浸式体验设计研究[J].艺术百家,2018,(04):161-169.',
    '[10]李宣,舒安琦.交互叙事视角下智慧博物馆沉浸式体验设计研究[J].包装工程,2024,45(06):461-470.'
    'DOI:10.19554/j.cnki.1001-3563.2024.06.052.',
    '[11]范浩宇,蔡新元.智慧博物馆沉浸式体验空间营造研究[J].家具与室内装饰,2023,30(09):117-123.'
    'DOI:10.16771/j.cn43-1247/ts.2023.09.018.',
    '[12]耿国华,高健,汤汶,等.文化遗产数字化保护与应用研究综述[J].西北大学学报(自然科学版),2025,55(01):1-22.'
    'DOI:10.16152/j.cnki.xdxbzr.2025-01-001.',
    '[13]李洪亮,杨和平.曾侯乙编钟演奏与创编虚拟仿真实验设计与实现[J].实验室研究与探索,2022,41(10):199-205.'
    'DOI:10.19927/j.cnki.syyt.2022.10.040.',
    '[14]彭红,张薇.湖北省博物馆IP开发的问题和策略[J].包装工程,2019,40(18):254-258.'
    'DOI:10.19554/j.cnki.1001-3563.2019.18.042.',
    '[15]Wojciechowski R,Walczak K,White M,et al.Building virtual and augmented reality museum '
    'exhibitions[C]//Proceedings of the Ninth International Conference on 3D Web Technology.'
    'New York:ACM,2004:135-144.DOI:10.1145/985040.985060.',
    '[16]Sylaiou S,Mania K,Karoulis A,et al.Exploring the relationship between presence and '
    'enjoyment in a virtual museum[J].International Journal of Human-Computer Studies,'
    '2010,68(05):243-253.DOI:10.1016/j.ijhcs.2009.11.002.',
    '[17]Bruno F,Bruno S,De Sensi G,et al.From 3D reconstruction to virtual reality:A complete '
    'methodology for digital archaeological exhibition[J].Journal of Cultural Heritage,'
    '2010,11(01):42-49.DOI:10.1016/j.culher.2009.02.006.',
]

# 段落文本替换（按前缀匹配旧段落）
REWRITE = [
    ('（一）对湖北省博物馆代表性馆藏的历史背景',
     '（一）对湖北省博物馆代表性馆藏的历史背景、器物形制、纹饰特征与色彩体系进行系统梳理与归纳[14]，'
     '建立清晰的展品视觉档案，为虚拟展厅的空间营造、展品建模与氛围设计提供依据与素材库。'),
    ('（一）理论研究：国内学术界对虚拟博物馆',
     '（一）理论研究：国内学者对虚拟博物馆与数字展陈的研究大致沿两条线索展开。一是技术实现线索：'
     '早期工作围绕三维场景描述语言与场景组织方式展开[3]，随后转向浏览器端实时渲染下的场景构造与交互漫游[4]、'
     'WebGL 环境中的碰撞检测[5]，以及基于 Unity3D 的信息可视化[6]与基于 Pano2VR 的全景漫游[7]等实现方案。'
     '二是设计研究线索：研究者从视觉认知与界面布局美度[8]、沉浸式体验的空间营造[9-11]等角度讨论数字展陈的'
     '观看经验，普遍认为虚拟展厅的核心价值不只是“把展品搬到网上”，而在于重建观看的情境与叙事逻辑；'
     '同时指出当前部分作品存在重技术轻内容、重还原轻体验、界面与交互割裂等问题，'
     '强调交互方式必须与内容逻辑相匹配。在文化遗产数字化方面，已有综述系统梳理了从数字化采集、三维重建、'
     '知识组织到展示应用的完整技术链条[12]。'),
    ('（二）设计实践：国内故宫博物院',
     '（二）设计实践：国内故宫博物院、敦煌研究院、上海博物馆等机构先后推出了数字全景漫游、'
     '高精度三维文物浏览与线上特展，形成了较为成熟的呈现方式，在文物数字化采集与虚拟展示方面积累了丰富经验；'
     '针对具体馆藏的数字化实践也已出现，如编钟类乐器的虚拟仿真演奏与创编实验[13]、'
     '以及围绕湖北省博物馆馆藏 IP 的开发研究[14]。但整体来看，知名大馆的数字化程度较高，'
     '地方性专题展陈与中小博物馆的虚拟化仍然薄弱，多采用全景照片拼接的方式呈现，'
     '缺少可自由漫游的真实三维空间，也缺少与馆藏内容深度结合的互动环节。'),
    ('国外对虚拟博物馆的研究与实践起步较早',
     '国外对虚拟博物馆的研究与实践起步较早，卢浮宫、大英博物馆、史密森学会等机构较早开展了藏品数字化与'
     '虚拟展厅建设，在元数据组织、展示规范与数字资源长期保存方面积累了经验，'
     '相关理论与馆建经验已有较为系统的梳理[1]。在工程实现方面，有研究提出了基于 Web3D 技术构建虚拟与'
     '增强现实博物馆展览的总体方案[15]，以及从三维重建到虚拟展示的完整方法论[17]；'
     '在体验研究方面，有研究通过实验考察了虚拟环境中“临场感”与体验愉悦度之间的关系，'
     '认为空间感、可操控性与内容深度共同影响用户的沉浸体验[16]。'
     '这些成果为本课题在渲染流程、交互方式与体验评价方面提供了参考。总体来看，'
     '国外研究在用户体验测评与研究方法上较为成熟，但针对中国地域文化与中国具体馆藏的空间叙事研究较少，'
     '面向楚文化的虚拟展厅设计仍有较大的探索空间。'),
]


def make_par(tmpl_p, text):
    newp = copy.deepcopy(tmpl_p)
    rs = newp.findall(qn('w:r'))
    if not rs:
        r = OxmlElement('w:r')
        newp.append(r)
        rs = [r]
    first = rs[0]
    for r in rs[1:]:
        newp.remove(r)
    for ch in list(first):
        if ch.tag != qn('w:rPr'):
            first.remove(ch)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    first.append(t)
    return newp


def set_text(p, text):
    rs = p.runs
    if not rs:
        p.add_run(text)
        return
    rs[0].text = text
    for r in rs[1:]:
        r._element.getparent().remove(r._element)


def is_ref(t):
    t = t.strip()
    return t.startswith('[') and len(t) > 2 and t[1].isdigit() and ']' in t[:4]


def patch_proposal():
    d = Document(os.path.join(HERE, '开题报告.docx'))
    t = d.tables[1]
    changed = 0
    for ri in range(len(t.rows)):
        cell = t.rows[ri].cells[0]
        ps = cell.paragraphs
        # 参考文献
        if ps and ps[0].text.strip().startswith('已查阅的主要参考文献'):
            tmpl = None
            for p in ps[1:]:
                if is_ref(p.text):
                    tmpl = p._p
                    break
            if tmpl is None:
                raise SystemExit('找不到参考文献模板段落')
            for p in ps[1:]:
                p._p.getparent().remove(p._p)
            prev = ps[0]._p
            for r in REFS:
                newp = make_par(tmpl, r)
                prev.addnext(newp)
                prev = newp
            changed += 1
            continue
        # 正文段落替换
        for p in ps:
            for old_prefix, new_text in REWRITE:
                if p.text.strip().startswith(old_prefix):
                    set_text(p, new_text)
                    changed += 1
    out = os.path.join(OUT, '开题报告.docx')
    d.save(out)
    print('开题报告: 修改 %d 处 ->' % changed, out)


def patch_taskbook():
    d = Document(os.path.join(HERE, '任务书.docx'))
    refs = [p for p in d.paragraphs if is_ref(p.text)]
    if not refs:
        raise SystemExit('任务书里找不到参考文献段落')
    tmpl = refs[0]._p
    n_old = len(refs)
    # 前 n_old 段就地改写，多出来的克隆插入
    prev = None
    for i, p in enumerate(refs):
        if i < len(REFS):
            set_text(p, REFS[i])
            prev = p._p
        else:
            p._p.getparent().remove(p._p)
    if len(REFS) > n_old:
        for r in REFS[n_old:]:
            newp = make_par(tmpl, r)
            prev.addnext(newp)
            prev = newp
    out = os.path.join(OUT, '任务书.docx')
    d.save(out)
    print('任务书: 参考文献 %d -> %d 条 ->' % (n_old, len(REFS)), out)


if __name__ == '__main__':
    patch_proposal()
    patch_taskbook()
